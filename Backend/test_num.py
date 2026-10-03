from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
p = doc.add_paragraph('Item 1')
p.style = 'List Number'

p2 = doc.add_paragraph('Item 2')
p2.style = 'List Number'

# Reset numbering for Item 1 of next list?
p3 = doc.add_paragraph('Item 1 of List 2')
p3.style = 'List Number'

doc.save('d:/neomakalah/test_num.docx')
