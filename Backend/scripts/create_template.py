from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path


def create_uin_template():
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    p = doc.add_paragraph()
    run = p.add_run("UNIVERSITAS ISLAM NEGERI ALAUDDIN MAKASSAR")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("{{fakultas}}")
    run.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("{{jurusan}}")
    run.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("MAKALAH")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("{{mata_kuliah}}")
    run.bold = True
    run.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("{{judul}}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.underline = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Oleh:")
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("{{penulis_1}}")
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NIM: {{nim_penulis_1}}")

    if "{{penulis_2}}" in doc.text:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("{{penulis_2}}")
        run.bold = True

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("NIM: {{nim_penulis_2}}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("{{dosen}}")
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("{{kampus}}")
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("{{tahun}}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    doc.add_page_break()

    p = doc.add_paragraph()
    run = p.add_run("ABSTRAK")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("{{abstrak}}")

    doc.add_page_break()

    p = doc.add_paragraph()
    run = p.add_run("BAB I")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("PENDAHULUAN")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("1.1 Latar Belakang")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("{{latar_belakang}}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH_LEFT
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("1.2 Rumusan Masalah")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("{{rumusan_masalah}}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH_LEFT
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("1.3 Tujuan")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("{{tujuan}}")

    doc.add_page_break()

    p = doc.add_paragraph()
    run = p.add_run("BAB II")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("PEMBAHASAN")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("{{pembahasan}}")

    doc.add_page_break()

    p = doc.add_paragraph()
    run = p.add_run("BAB III")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("PENUTUP")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("3.1 Kesimpulan")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("{{kesimpulan}}")

    doc.add_page_break()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DAFTAR PUSTAKA")
    run.bold = True
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n")
    run.font.size = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("{{tempat_tanggal}}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH_RIGHT
    run = p.add_run("Penulis")

    templates_dir = Path("templates")
    templates_dir.mkdir(parents=True, exist_ok=True)

    template_path = templates_dir / "template_uin.docx"
    doc.save(str(template_path))

    print(f"✓ Template berhasil dibuat: {template_path}")
    return template_path


if __name__ == "__main__":
    create_uin_template()