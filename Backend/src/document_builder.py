import re
from pathlib import Path
from datetime import datetime
from typing import Dict

from docx import Document
from rich.console import Console


class DocumentBuilder:
    def __init__(self, output_dir="output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.console = Console()

    def build_document(self, template: Document, content: Dict[str, str]) -> Document:
        from src.template_engine import TemplateEngine


        engine = TemplateEngine()
        document, replacement_status = engine.replace_placeholders(template, content)

        replaced_count = sum(1 for status in replacement_status.values() if status)
        total_count = len(replacement_status)
        accuracy = (replaced_count / total_count * 100) if total_count > 0 else 0

        self.console.print(
            f"[cyan][OK] Placeholder replacement: {replaced_count}/{total_count} ({accuracy:.1f}%)[/cyan]"
        )
        if accuracy < 98:
            self.console.print("[yellow][WARNING] Placeholder replacement di bawah 98%[/yellow]")

        return document

    def save_document(self, document: Document, filename: str) -> str:
        output_path = self.output_dir / filename
        document.save(output_path)
        return str(output_path)

    def generate_filename(self, metadata: Dict[str, str]) -> str:
        judul = metadata.get("judul", "makalah").lower()
        judul = " ".join([word for word in judul.split() if word.isalnum()])
        judul = judul[:50]

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{judul}_{timestamp}.docx"

    def format_section(self, text: str, section_type: str) -> str:
        """Clean and normalize section text before template insertion."""
        if not text:
            return text

        formatted_text = text.strip()

        # Strip any remaining markdown artifacts (safety net)
        # Remove bold markers
        formatted_text = re.sub(r'\*\*(.+?)\*\*', r'\1', formatted_text)
        formatted_text = re.sub(r'__(.+?)__', r'\1', formatted_text)
        # Remove heading markers
        formatted_text = re.sub(r'^#{1,6}\s+', '', formatted_text, flags=re.MULTILINE)
        # Remove markdown bullets
        formatted_text = re.sub(r'^[\-\*]\s+', '', formatted_text, flags=re.MULTILINE)

        # Normalize line endings
        formatted_text = formatted_text.replace('\r\n', '\n').replace('\r', '\n')

        # Collapse 3+ consecutive newlines into 2 (paragraph separation)
        formatted_text = re.sub(r'\n{3,}', '\n\n', formatted_text)

        if section_type == "judul":
            formatted_text = formatted_text.title()
        elif section_type in ["latar_belakang", "kesimpulan", "pembahasan"]:
            # Ensure each paragraph ends with a period
            paragraphs = formatted_text.split('\n\n')
            cleaned = []
            for para in paragraphs:
                para = para.strip()
                if para and not para.endswith('.') and not para.endswith('?') and not para.endswith('!'):
                    para += '.'
                cleaned.append(para)
            formatted_text = '\n\n'.join(cleaned)

        return formatted_text

    def validate_document(self, document: Document) -> bool:
        try:
            if not document.paragraphs:
                return False
            return any(p.text.strip() for p in document.paragraphs)
        except Exception as exc:
            self.console.print(f"[red][FAIL] Error validating document: {exc}[/red]")
            return False

    def display_completion_summary(self, output_path: str, token_usage: Dict[str, int]):
        self.console.print("\n[bold green]MAKALAH BERHASIL DIBUAT[/bold green]\n")
        self.console.print(f"[cyan]File:[/cyan] {output_path}")
        self.console.print(
            f"[white]Input tokens: {token_usage.get('total_input_tokens', 0)} | "
            f"Output tokens: {token_usage.get('total_output_tokens', 0)} | "
            f"Total: {token_usage.get('total_tokens', 0)}[/white]"
        )

    def build_complete_paper_data(
        self,
        metadata: Dict[str, str],
        authors: list,
        institution: Dict[str, str],
        content: Dict[str, str],
    ) -> Dict[str, str]:
        paper_data = {}
        paper_data.update(metadata)

        for i, author in enumerate(authors, 1):
            paper_data[f"penulis_{i}"] = author.get("nama", "")
            paper_data[f"nim_penulis_{i}"] = author.get("nim", "")

        # Zero out remaining placeholders up to 5
        for i in range(len(authors) + 1, 6):
            paper_data[f"penulis_{i}"] = ""
            paper_data[f"nim_penulis_{i}"] = ""

        paper_data.update(institution)
        paper_data.update(content)

        # Parse rumusan_masalah (Pisahkan intro dan list angka)
        rumusan_text = content.get("rumusan_masalah", "")
        intro_r, items_r = [], []
        for line in rumusan_text.split('\n'):
            line = line.strip()
            if not line: continue
            if re.match(r'^\d+[\.\)]\s*', line):
                items_r.append(re.sub(r'^\d+[\.\)]\s*', '', line).strip())
            else:
                if items_r:
                    items_r[-1] += " " + line
                else:
                    intro_r.append(line)
        paper_data["rumusan_masalah"] = '\n'.join(intro_r)
        for i in range(1, 6):
            # Back to clean text, let Word handle the numbering
            paper_data[f"rumusan{i}"] = items_r[i-1] if i-1 < len(items_r) else ""

        # Parse tujuan (Pisahkan intro dan list angka)
        tujuan_text = content.get("tujuan", "")
        intro_t, items_t = [], []
        for line in tujuan_text.split('\n'):
            line = line.strip()
            if not line: continue
            if re.match(r'^\d+[\.\)]\s*', line):
                items_t.append(re.sub(r'^\d+[\.\)]\s*', '', line).strip())
            else:
                if items_t:
                    items_t[-1] += " " + line
                else:
                    intro_t.append(line)
        paper_data["tujuan"] = '\n'.join(intro_t)
        for i in range(1, 6):
            # Back to clean text, let Word handle the numbering
            paper_data[f"tujuan{i}"] = items_t[i-1] if i-1 < len(items_t) else ""

        # Parse pembahasan (Pisahkan heading A. B. C. D. dan isinya)
        pembahasan_text = content.get("pembahasan", "")
        pemb_headings, pemb_contents = [], []
        curr_heading, curr_content = "", []
        
        for line in pembahasan_text.split('\n'):
            line = line.strip()
            if not line: continue
            
            # Detect heading like "A. Definisi"
            if re.match(r'^[A-Z][\.\)]\s+', line):
                if curr_heading:
                    pemb_headings.append(curr_heading)
                    pemb_contents.append(curr_content)
                
                # Strip the prefix "A. " out of the heading
                # So if the AI gives "A. Definisi Sistem", it saves "Definisi Sistem"
                curr_heading = re.sub(r'^[A-Z][\.\)]\s+', '', line).strip()
                curr_content = []
            else:
                if curr_heading:
                    curr_content.append(line)
        
        # Save the last item
        if curr_heading:
            pemb_headings.append(curr_heading)
            pemb_contents.append(curr_content)
            
        # Assign defaults if missing and map to isiX.1, isiX.2, isiX.3
        for i in range(1, 6):
            if i - 1 < len(pemb_headings):
                paper_data[f"pembahasan{i}"] = pemb_headings[i-1]
                paras = pemb_contents[i-1]
                # Map full content for single-placeholder templates (like {{isi1}})
                paper_data[f"isi{i}"] = '\n\n'.join(paras)
                # Map per-paragraph for multi-placeholder templates (like {{isi1.1}})
                for j in range(1, 4):
                    paper_data[f"isi{i}.{j}"] = paras[j-1] if j - 1 < len(paras) else ""
            else:
                paper_data[f"pembahasan{i}"] = f"Pembahasan {i}"
                paper_data[f"isi{i}"] = ""
                for j in range(1, 4):
                    paper_data[f"isi{i}.{j}"] = ""

        # Parse Kesimpulan (Pisahkan menjadi 2 paragraf)
        kesimpulan_text = content.get("kesimpulan", "")
        kesimpulan_paras = [p.strip() for p in kesimpulan_text.split('\n') if p.strip()]
        paper_data["kesimpulan1.1"] = kesimpulan_paras[0] if len(kesimpulan_paras) > 0 else "Kesimpulan belum tersedia."
        paper_data["kesimpulan1.2"] = kesimpulan_paras[1] if len(kesimpulan_paras) > 1 else ""

        abstrak = content.get("abstrak", "").strip()
        if not abstrak:
            abstrak = content.get("latar_belakang", "")

        if abstrak and len(abstrak) > 300:
            paper_data["abstrak"] = abstrak[:300] + "..."
        else:
            paper_data["abstrak"] = abstrak

        tempat = institution.get("tempat_pembuatan", "").strip()
        tanggal = institution.get("tanggal_pembuatan", "").strip()
        if tempat and tanggal:
            paper_data["tempat_tanggal"] = f"{tempat}, {tanggal}"
        elif tempat:
            paper_data["tempat_tanggal"] = tempat
        elif tanggal:
            paper_data["tempat_tanggal"] = tanggal

        optional_fields = [
            "penulis_2",
            "nim_penulis_2",
            "kata_pengantar",
            "tempat_pembuatan",
            "tanggal_pembuatan",
        ]
        for field in optional_fields:
            paper_data.setdefault(field, "")

        # Clean string "DAFTAR PUSTAKA" prefix from AI output if present
        if "daftar_pustaka" in paper_data:
            dp_text = paper_data["daftar_pustaka"].strip()
            dp_text = re.sub(r'(?i)^\**daftar\s+pustaka\**\s*\n*', '', dp_text).strip()
            paper_data["daftar_pustaka"] = dp_text
        else:
            paper_data["daftar_pustaka"] = ""

        return paper_data


def get_document_builder(output_dir="output"):
    return DocumentBuilder(output_dir)
